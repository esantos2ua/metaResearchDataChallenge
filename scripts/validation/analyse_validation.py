"""Analyse the returned B2 validation coding: completeness, agreement, adjudication, precision.

Three steps, run in order as the sheets come back:

    python scripts/validation/analyse_validation.py check
        Which coder sheets are in `returned/` (and adjudication sheets in `adjudication/returned/`),
        and is each one complete and untampered?

    python scripts/validation/analyse_validation.py adjudicate
        Cohen's κ (pooled and per pair) and Fleiss' κ for both questions, plus Gwet's AC1 and
        per-category specific agreement (supplementary, not preregistered), then one blind
        adjudication workbook per adjudicator listing only the records their pair disagreed on.

    python scripts/validation/analyse_validation.py results [--provisional]
        Final label per record (agreed code, else the adjudicator's), precision with Wilson 95% CIs,
        exploratory breakdowns by era / language / primary topic, and the publishable coded sample.

Rules fixed in docs/PROTOCOL.md (B2) and docs/CODING_GUIDE.md:
- Any mismatch between the two coders (including YES vs UNCLEAR) goes to adjudication.
- Adjudicators work blind: they see the record, not the coders' answers or notes.
- Primary precision counts a final UNCLEAR as out of scope; the sensitivity estimate drops
  UNCLEAR from the denominator.
- Breakdowns by era, language and topic are exploratory.
Coders appear in the published sample under pseudonymous IDs (C1–C8), not by name.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import re
import sys
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from make_coder_sheets import COLUMNS, PAIRS, slug  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
VALIDATION = ROOT / "data" / "validation"
MASTER_PATH = VALIDATION / "validation_sample_200_master.json"
RECORDS_PATH = ROOT / "dashboard" / "app" / "data" / "records.json"
QUERY_CONFIG = ROOT / "scripts" / "extraction" / "query_config.yaml"

# Third coder per pair: a co-lead who is not in that pair (TEAM-ROSTER.md).
ADJUDICATORS = {
    "pair1_bruno_elena": "Eduardo Santos",
    "pair2_stephanie_valentin": "Shinichi Nakagawa",
    "pair3_shinichi_losia": "Marija Purgar",
    "pair4_marija_eduardo": "Losia Lagisz",
}
CODER_IDS = {name: f"C{i}" for i, name in
             enumerate((c for _, _, coders in PAIRS for c in coders), start=1)}

VALID = ("YES", "NO", "UNCLEAR")
QUESTIONS = {"q1": "is_metaresearch", "q2": "is_canadian"}
Z95 = 1.959964

ANSWER_FILL = PatternFill("solid", fgColor="FFF6D5")
HEADER_FILL = PatternFill("solid", fgColor="DDE5F0")


# ---------------------------------------------------------------- statistics

def wilson(k: int, n: int, z: float = Z95) -> tuple[float, float, float]:
    """Proportion k/n with its Wilson score interval."""
    if n == 0:
        return (math.nan, math.nan, math.nan)
    p = k / n
    denom = 1 + z * z / n
    centre = (p + z * z / (2 * n)) / denom
    half = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / denom
    return (p, max(0.0, centre - half), min(1.0, centre + half))


def cohen_kappa(a: list[str], b: list[str]) -> float:
    n = len(a)
    if n == 0:
        return math.nan
    po = sum(x == y for x, y in zip(a, b)) / n
    ca, cb = Counter(a), Counter(b)
    pe = sum(ca[c] * cb[c] for c in set(a) | set(b)) / (n * n)
    return math.nan if pe == 1 else (po - pe) / (1 - pe)


def gwet_ac1(a: list[str], b: list[str], categories: tuple[str, ...] = VALID) -> float:
    """Gwet's AC1: chance agreement from category spread, so it stays stable when one answer dominates.

    Not preregistered (κ is); reported alongside κ because Q2 is nearly all YES (kappa paradox).
    """
    n, q = len(a), len(categories)
    if n == 0:
        return math.nan
    po = sum(x == y for x, y in zip(a, b)) / n
    pi = [(a.count(c) + b.count(c)) / (2 * n) for c in categories]
    pe = sum(p * (1 - p) for p in pi) / (q - 1)
    return math.nan if pe == 1 else (po - pe) / (1 - pe)


def specific_agreement(a: list[str], b: list[str]) -> dict[str, float]:
    """Per-category specific agreement 2·n_kk / (n_k(a) + n_k(b)); e.g. YES = positive agreement."""
    out = {}
    for c in VALID:
        denom = a.count(c) + b.count(c)
        both = sum(x == y == c for x, y in zip(a, b))
        out[c] = 2 * both / denom if denom else math.nan
    return out


def fleiss_kappa(pairs: list[tuple[str, str]]) -> float:
    """Fleiss' κ for two ratings per record, raters not fixed across records."""
    n = len(pairs)
    if n == 0:
        return math.nan
    cats = sorted({c for p in pairs for c in p})
    pj = {c: sum(p.count(c) for p in pairs) / (2 * n) for c in cats}
    p_bar = sum(1.0 if x == y else 0.0 for x, y in pairs) / n  # P_i for m=2 raters
    pe = sum(v * v for v in pj.values())
    return math.nan if pe == 1 else (p_bar - pe) / (1 - pe)


# ---------------------------------------------------------------- loading

def corpus_topics() -> dict[str, str]:
    """{topic_id: name} for the corpus topics in query_config.yaml (read without a YAML dependency)."""
    text = QUERY_CONFIG.read_text(encoding="utf-8")
    return dict(re.findall(r"-\s*id:\s*(T\d+)\s*\n\s*name:\s*(.+)", text))


def load_master() -> dict[int, dict]:
    master = {int(r["record_id"]): r for r in json.loads(MASTER_PATH.read_text(encoding="utf-8"))}
    langs = {}
    if RECORDS_PATH.exists():
        langs = {r["id"]: r.get("language") or "unknown"
                 for r in json.loads(RECORDS_PATH.read_text(encoding="utf-8"))["records"]}
    for r in master.values():
        r["language"] = langs.get(r["openalex_id"], "unknown")
    return master


def expected_sheets() -> list[dict]:
    out = []
    for pair_id, rng, coders in PAIRS:
        start, end = (int(x) for x in rng.split("_"))
        for coder in coders:
            out.append({"pair": pair_id, "coder": coder, "ids": list(range(start, end + 1)),
                        "file": f"coding_sheet_{slug(coder)}_records_{rng}.xlsx"})
    return out


def norm(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip().upper()
    return s or None


def read_sheet(path: Path, sheet: str, fields: list[str]) -> list[dict]:
    ws = load_workbook(path, read_only=True, data_only=True)[sheet]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    out = []
    for row in rows:
        if all(v is None for v in row):
            continue
        out.append({f: row[header.index(h)] if h in header else None for f, h in fields})
    return out


CODER_FIELDS = [(f, h) for f, h, _ in COLUMNS]


def validate_coder_sheet(path: Path, spec: dict, master: dict[int, dict]) -> tuple[list[dict], list[str]]:
    problems = []
    rows = read_sheet(path, "Coding", CODER_FIELDS)
    ids = [int(r["record_id"]) if r["record_id"] is not None else None for r in rows]
    if ids != spec["ids"]:
        problems.append(f"record_ids changed (expected {spec['ids'][0]}–{spec['ids'][-1]} in order)")
    for r in rows:
        rid = r["record_id"]
        if rid is None or int(rid) not in master:
            continue
        if (r["title"] or "") != (master[int(rid)]["title"] or ""):
            problems.append(f"record {rid}: title differs from master (row moved or edited?)")
        for q, field in QUESTIONS.items():
            v = norm(r[field])
            if v is None:
                problems.append(f"record {rid}: {field} blank")
            elif v not in VALID:
                problems.append(f"record {rid}: {field} = {r[field]!r} (not YES/NO/UNCLEAR)")
            r[field] = v
    return rows, problems


def load_codings(returned: Path, master: dict[int, dict], strict: bool) -> dict[str, dict]:
    """{pair_id: {"coders": [a, b], "rows": {coder: {record_id: row}}}}; exits on problems if strict."""
    pairs: dict[str, dict] = defaultdict(lambda: {"coders": [], "rows": {}})
    missing, bad = [], {}
    for spec in expected_sheets():
        path = returned / spec["file"]
        pairs[spec["pair"]]["coders"].append(spec["coder"])
        if not path.exists():
            missing.append(spec)
            continue
        rows, problems = validate_coder_sheet(path, spec, master)
        if problems:
            bad[spec["coder"]] = problems
        pairs[spec["pair"]]["rows"][spec["coder"]] = {int(r["record_id"]): r for r in rows
                                                      if r["record_id"] is not None}
    if strict and (missing or bad):
        for s in missing:
            print(f"  missing: {s['coder']} ({s['file']})", file=sys.stderr)
        for coder, probs in bad.items():
            print(f"  {coder}: {len(probs)} problem(s), first: {probs[0]}", file=sys.stderr)
        sys.exit("Not all coder sheets are in and valid; run `check` for details.")
    return dict(pairs)


# ---------------------------------------------------------------- check

def cmd_check(args) -> None:
    master = load_master()
    n_ok = 0
    for spec in expected_sheets():
        path = args.returned / spec["file"]
        if not path.exists():
            print(f"[ ] {spec['coder']:<18} not returned ({spec['file']})")
            continue
        _, problems = validate_coder_sheet(path, spec, master)
        if problems:
            print(f"[!] {spec['coder']:<18} {len(problems)} problem(s):")
            for p in problems[:10]:
                print(f"      - {p}")
            if len(problems) > 10:
                print(f"      … and {len(problems) - 10} more")
        else:
            n_ok += 1
            print(f"[x] {spec['coder']:<18} complete")
    print(f"\n{n_ok}/8 sheets complete.")
    for pair_id, adj in ADJUDICATORS.items():
        f = args.adjudicated / f"adjudication_{slug(adj)}.xlsx"
        if f.exists():
            print(f"Adjudication from {adj} ({pair_id}) is in.")


# ---------------------------------------------------------------- adjudicate

def pair_disagreements(pair: dict) -> dict[int, list[str]]:
    a, b = pair["coders"]
    out = {}
    for rid in sorted(pair["rows"][a]):
        qs = [q for q, f in QUESTIONS.items() if pair["rows"][a][rid][f] != pair["rows"][b][rid][f]]
        if qs:
            out[rid] = qs
    return out


def agreement_stats(pairs: dict[str, dict]) -> dict:
    stats: dict = {"pooled": {}, "per_pair": {}}
    for q, field in QUESTIONS.items():
        all_pairs = []
        for pair_id, pair in pairs.items():
            a, b = pair["coders"]
            ra = [pair["rows"][a][rid][field] for rid in sorted(pair["rows"][a])]
            rb = [pair["rows"][b][rid][field] for rid in sorted(pair["rows"][a])]
            all_pairs += list(zip(ra, rb))
            stats["per_pair"].setdefault(pair_id, {})[q] = {
                "n": len(ra), "agreement": sum(x == y for x, y in zip(ra, rb)) / len(ra),
                "cohen_kappa": cohen_kappa(ra, rb), "gwet_ac1": gwet_ac1(ra, rb)}
        pa, pb = [x for x, _ in all_pairs], [y for _, y in all_pairs]
        stats["pooled"][q] = {
            "n": len(all_pairs),
            "agreement": sum(x == y for x, y in all_pairs) / len(all_pairs),
            "cohen_kappa": cohen_kappa(pa, pb),
            "fleiss_kappa": fleiss_kappa(all_pairs),
            "gwet_ac1": gwet_ac1(pa, pb),
            "specific_agreement": specific_agreement(pa, pb),
            "table": Counter(f"{x}/{y}" for x, y in all_pairs).most_common()}
    return stats


def write_adjudication_sheet(path: Path, adjudicator: str, pair_id: str,
                             todo: dict[int, list[str]], master: dict[int, dict]) -> None:
    wb = Workbook()
    info = wb.active
    info.title = "Instructions"
    lines = [
        "Canadian Metaresearch Dashboard: adjudication (B2)",
        "",
        f"Adjudicator: {adjudicator}",
        f"Records: {len(todo)} on which the two coders of one pair disagreed",
        "",
        "1. Apply docs/CODING_GUIDE.md (v1.0), exactly as the coders did.",
        "2. You are blind to the coders' answers on purpose: judge each record afresh.",
        "3. Answer only the question(s) marked NEEDED; leave the other answer cell empty.",
        "4. Do not look the record up in OpenAlex or on the dashboard.",
        "5. Your answer is final. UNCLEAR is allowed and counts as out of scope in the primary estimate.",
        f"6. Save as adjudication_{slug(adjudicator)}.xlsx in data/validation/adjudication/returned/.",
    ]
    for i, line in enumerate(lines, start=1):
        c = info.cell(row=i, column=1, value=line)
        if i == 1:
            c.font = Font(bold=True, size=14)
    info.column_dimensions["A"].width = 110

    ws = wb.create_sheet("Adjudication")
    wb.active = 1
    cols = [("record_id", 9), ("year", 7), ("venue", 24), ("title", 45), ("abstract", 80),
            ("authors_and_affiliations", 45), ("doi", 20),
            ("Q1 needed", 10), ("adj_is_metaresearch", 14),
            ("Q2 needed", 10), ("adj_is_canadian", 14), ("notes", 35)]
    for j, (h, w) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=j, value=h)
        c.font, c.fill = Font(bold=True), HEADER_FILL
        ws.column_dimensions[c.column_letter].width = w
    for i, (rid, qs) in enumerate(sorted(todo.items()), start=2):
        m = master[rid]
        values = [rid, int(m["year"]) if m.get("year") else None, m.get("venue"), m.get("title"),
                  m.get("abstract"), m.get("authors_and_affiliations"), m.get("doi"),
                  "NEEDED" if "q1" in qs else "—", None,
                  "NEEDED" if "q2" in qs else "—", None, None]
        for j, v in enumerate(values, start=1):
            c = ws.cell(row=i, column=j, value=v or None)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for j, q in ((9, "q1"), (11, "q2")):
            if q in qs:
                ws.cell(row=i, column=j).fill = ANSWER_FILL
        ws.cell(row=i, column=12).fill = ANSWER_FILL
    last = len(todo) + 1
    dv = DataValidation(type="list", formula1='"YES,NO,UNCLEAR"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"I2:I{max(last, 2)}")
    dv.add(f"K2:K{max(last, 2)}")
    ws.freeze_panes = "B2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def fmt(x: float) -> str:
    return "—" if x is None or (isinstance(x, float) and math.isnan(x)) else f"{x:.2f}"


def cmd_adjudicate(args) -> None:
    master = load_master()
    pairs = load_codings(args.returned, master, strict=True)
    stats = agreement_stats(pairs)
    print("Inter-coder agreement (200 records, 4 pairs)")
    for q, field in QUESTIONS.items():
        s = stats["pooled"][q]
        print(f"  {field:<16} agreement {s['agreement']:.1%}  Cohen κ {fmt(s['cohen_kappa'])}  "
              f"Fleiss κ {fmt(s['fleiss_kappa'])}  Gwet AC1 {fmt(s['gwet_ac1'])}  (target κ ≥ 0.70)")
        sa = s["specific_agreement"]
        print(f"  {'':<16} specific agreement YES {fmt(sa['YES'])}  NO {fmt(sa['NO'])}  "
              f"UNCLEAR {fmt(sa['UNCLEAR'])}")
    for pair_id, s in stats["per_pair"].items():
        print(f"  {pair_id:<26} κ Q1 {fmt(s['q1']['cohen_kappa'])}  κ Q2 {fmt(s['q2']['cohen_kappa'])}  "
              f"AC1 Q1 {fmt(s['q1']['gwet_ac1'])}  AC1 Q2 {fmt(s['q2']['gwet_ac1'])}")

    out_dir = args.out / "adjudication"
    total = 0
    for pair_id, pair in pairs.items():
        adj = ADJUDICATORS[pair_id]
        assert adj not in pair["coders"], f"{adj} cannot adjudicate their own pair"
        todo = pair_disagreements(pair)
        total += len(todo)
        path = out_dir / f"adjudication_{slug(adj)}.xlsx"
        write_adjudication_sheet(path, adj, pair_id, todo, master)
        print(f"  {adj:<18} {len(todo):>3} record(s) from {pair_id} -> {path}")
    print(f"\n{total} record(s) to adjudicate.")
    (args.out / "validation_agreement.json").write_text(json.dumps(stats, indent=2), encoding="utf-8")


# ---------------------------------------------------------------- results

def load_adjudications(returned: Path) -> dict[int, dict]:
    fields = [("record_id", "record_id"), ("q1", "adj_is_metaresearch"),
              ("q2", "adj_is_canadian"), ("notes", "notes")]
    out = {}
    for pair_id, adj in ADJUDICATORS.items():
        path = returned / f"adjudication_{slug(adj)}.xlsx"
        if not path.exists():
            continue
        for r in read_sheet(path, "Adjudication", fields):
            if r["record_id"] is not None:
                out[int(r["record_id"])] = {"adjudicator": adj, "q1": norm(r["q1"]),
                                            "q2": norm(r["q2"]), "notes": r["notes"]}
    return out


def precision_block(finals: list[str]) -> dict:
    c = Counter(finals)
    n = len(finals)
    p, lo, hi = wilson(c["YES"], n)
    ps, los, his = wilson(c["YES"], c["YES"] + c["NO"])
    return {"n": n, "yes": c["YES"], "no": c["NO"], "unclear": c["UNCLEAR"],
            "primary": {"estimate": p, "ci95": [lo, hi]},
            "unclear_excluded": {"n": c["YES"] + c["NO"], "estimate": ps, "ci95": [los, his]}}


def cmd_results(args) -> None:
    master = load_master()
    pairs = load_codings(args.returned, master, strict=True)
    adjud = load_adjudications(args.adjudicated)
    topics = corpus_topics()
    topic_names = {n.strip().lower() for n in topics.values()}
    stats = agreement_stats(pairs)

    rows, unresolved = [], []
    for pair_id, pair in pairs.items():
        a, b = pair["coders"]
        for rid in sorted(pair["rows"][a]):
            ra, rb, m = pair["rows"][a][rid], pair["rows"][b][rid], master[rid]
            row = {"record_id": rid, "openalex_id": m["openalex_id"], "doi": m.get("doi"),
                   "year": m.get("year"), "stratum_era": m.get("stratum_era"),
                   "language": m["language"], "primary_topic": m.get("primary_topic"),
                   "assigned_topics": m.get("assigned_topics"),
                   "entry": ("primary topic" if (m.get("primary_topic") or "").lower() in topic_names
                             else "secondary topic only"),
                   "coder_a": CODER_IDS[a], "coder_b": CODER_IDS[b]}
            adj = adjud.get(rid)
            for q, field in QUESTIONS.items():
                row[f"{field}_a"], row[f"{field}_b"] = ra[field], rb[field]
                if ra[field] == rb[field]:
                    final, how = ra[field], "agreed"
                elif adj and adj[q] in VALID:
                    final, how = adj[q], "adjudicated"
                else:
                    final, how = "UNCLEAR", "unresolved"
                    unresolved.append((rid, field))
                row[f"{field}_adjudicated"] = adj[q] if adj and how == "adjudicated" else None
                row[f"{field}_final"], row[f"{field}_resolution"] = final, how
            row["adjudicator"] = (CODER_IDS.get(adj["adjudicator"]) if adj else None)
            row["notes_a"], row["notes_b"] = ra.get("notes"), rb.get("notes")
            row["notes_adjudicator"] = adj.get("notes") if adj else None
            rows.append(row)

    if unresolved and not args.provisional:
        sys.exit(f"{len(unresolved)} disagreement(s) not yet adjudicated, e.g. record "
                 f"{unresolved[0][0]} ({unresolved[0][1]}). Wait for the adjudication sheets, "
                 f"or rerun with --provisional (unresolved counted as UNCLEAR).")

    results = {"generated": date.today().isoformat(), "provisional": bool(unresolved),
               "unresolved": len(unresolved), "agreement": stats, "precision": {}, "exploratory": {}}
    for q, field in QUESTIONS.items():
        results["precision"][field] = precision_block([r[f"{field}_final"] for r in rows])
    # Provisional only: where precision can land once adjudication is in (every open
    # disagreement resolved NO vs. every one resolved YES).
    results["bounds"] = {}
    if unresolved:
        for field in QUESTIONS.values():
            yes = sum(r[f"{field}_final"] == "YES" for r in rows)
            open_ = sum(r[f"{field}_resolution"] == "unresolved" for r in rows)
            results["bounds"][field] = {"unresolved": open_, "min": yes / len(rows),
                                        "max": (yes + open_) / len(rows)}
    results["precision"]["both"] = precision_block(
        ["YES" if r["is_metaresearch_final"] == r["is_canadian_final"] == "YES" else
         "UNCLEAR" if "UNCLEAR" in (r["is_metaresearch_final"], r["is_canadian_final"]) else "NO"
         for r in rows])
    for key in ("stratum_era", "language", "entry"):
        groups = defaultdict(list)
        for r in rows:
            g = r[key] if key != "language" or r[key] in ("en", "fr", "unknown") else "other"
            groups[g or "unknown"].append(r["is_metaresearch_final"])
        results["exploratory"][key] = {g: precision_block(v) for g, v in sorted(groups.items())}
    # A record carrying several corpus topics counts once in each of them.
    by_topic = defaultdict(list)
    for r in rows:
        for tid in re.findall(r"T\d+", str(r["assigned_topics"] or "")):
            if tid in topics:
                by_topic[f"{tid} {topics[tid]}"].append(r["is_metaresearch_final"])
    results["exploratory"]["corpus_topic"] = {g: precision_block(v) for g, v in sorted(by_topic.items())}

    suffix = "_PROVISIONAL" if unresolved else ""
    stamp = date.today().isoformat()
    args.out.mkdir(parents=True, exist_ok=True)
    csv_path = args.out / f"relevance_sample_{stamp}{suffix}.csv"
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(sorted(rows, key=lambda r: r["record_id"]))
    json_path = args.out / f"validation_results{suffix}.json"
    json_path.write_text(json.dumps(results, indent=2), encoding="utf-8")
    md_path = args.out / f"validation_results{suffix}.md"
    md_path.write_text(render_markdown(results), encoding="utf-8")
    print(render_markdown(results))
    print(f"Wrote {csv_path}\nWrote {json_path}\nWrote {md_path}")


def ci(block: dict, which: str = "primary") -> str:
    e = block[which]
    return f"{fmt(e['estimate'])} (95% CI {fmt(e['ci95'][0])}–{fmt(e['ci95'][1])})"


def render_markdown(res: dict) -> str:
    L = ["# B2 validation results" + (" — PROVISIONAL" if res["provisional"] else ""), "",
         f"Generated {res['generated']}. Stratified random sample of 200 records from `corpus-v1`, "
         "every record double-coded blind to its OpenAlex topic; disagreements adjudicated blind.",
         ""]
    if res["provisional"]:
        L += [f"> **{res['unresolved']} disagreement(s) not yet adjudicated** and counted as UNCLEAR. "
              "Do not quote these numbers.", ""]
    L += ["## Precision", "",
          "| Question | YES / NO / UNCLEAR | Primary (UNCLEAR = out) | UNCLEAR excluded |",
          "| --- | --- | --- | --- |"]
    labels = {"is_metaresearch": "Q1 Metaresearch", "is_canadian": "Q2 Canadian",
              "both": "Q1 and Q2"}
    for key, label in labels.items():
        b = res["precision"][key]
        L.append(f"| {label} | {b['yes']} / {b['no']} / {b['unclear']} | {ci(b)} | "
                 f"{ci(b, 'unclear_excluded')} (n = {b['unclear_excluded']['n']}) |")
    if res.get("bounds"):
        L += ["", "Where the primary estimate can land once adjudication is in "
              "(point estimates, no CI):", "",
              "| Question | Open disagreements | All resolved NO | All resolved YES |",
              "| --- | --- | --- | --- |"]
        for field, b in res["bounds"].items():
            L.append(f"| {labels[field]} | {b['unresolved']} | {fmt(b['min'])} | {fmt(b['max'])} |")
    L += ["", "## Inter-coder agreement (before adjudication)", "",
          "| Question | Raw agreement | Cohen κ (pooled) | Fleiss κ | Gwet AC1 | "
          "Specific agreement YES / NO / UNCLEAR |", "| --- | --- | --- | --- | --- | --- |"]
    for q, field in QUESTIONS.items():
        s = res["agreement"]["pooled"][q]
        sa = s["specific_agreement"]
        L.append(f"| {field} | {s['agreement']:.1%} | {fmt(s['cohen_kappa'])} | {fmt(s['fleiss_kappa'])} | "
                 f"{fmt(s['gwet_ac1'])} | {fmt(sa['YES'])} / {fmt(sa['NO'])} / {fmt(sa['UNCLEAR'])} |")
    L += ["", "Cohen κ is the preregistered index (target ≥ 0.70). When one answer dominates "
          "(Q2 is nearly all YES), κ can be low despite high raw agreement (the kappa paradox), so "
          "Gwet's AC1 and per-category specific agreement are added as a **declared deviation**: "
          "supplementary, not a replacement for κ.", "",
          "| Pair | κ Q1 | κ Q2 | AC1 Q1 | AC1 Q2 | Agreement Q1 | Agreement Q2 |",
          "| --- | --- | --- | --- | --- | --- | --- |"]
    for pair_id, s in res["agreement"]["per_pair"].items():
        L.append(f"| {pair_id.split('_')[0]} | {fmt(s['q1']['cohen_kappa'])} | {fmt(s['q2']['cohen_kappa'])} | "
                 f"{fmt(s['q1']['gwet_ac1'])} | {fmt(s['q2']['gwet_ac1'])} | "
                 f"{s['q1']['agreement']:.0%} | {s['q2']['agreement']:.0%} |")
    titles = {"stratum_era": "era", "language": "language",
              "entry": "how the record entered the corpus",
              "corpus_topic": "corpus topic (a record with several topics counts in each)"}
    for key, title in titles.items():
        L += ["", f"## Exploratory: Q1 precision by {title}", "",
              "| Group | n | YES | Primary precision |", "| --- | --- | --- | --- |"]
        for g, b in res["exploratory"][key].items():
            L.append(f"| {g} | {b['n']} | {b['yes']} | {ci(b)} |")
    L += ["", "Subgroup intervals are wide by design; they are exploratory, not tests.", ""]
    return "\n".join(L)


# ---------------------------------------------------------------- main

def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("command", choices=("check", "adjudicate", "results"))
    ap.add_argument("--returned", type=Path, default=VALIDATION / "returned",
                    help="folder with the returned coder sheets")
    ap.add_argument("--adjudicated", type=Path, default=VALIDATION / "adjudication" / "returned",
                    help="folder with the returned adjudication sheets")
    ap.add_argument("--out", type=Path, default=VALIDATION, help="where outputs are written")
    ap.add_argument("--provisional", action="store_true",
                    help="results: count unadjudicated disagreements as UNCLEAR and mark outputs PROVISIONAL")
    args = ap.parse_args()
    {"check": cmd_check, "adjudicate": cmd_adjudicate, "results": cmd_results}[args.command](args)


if __name__ == "__main__":
    main()
