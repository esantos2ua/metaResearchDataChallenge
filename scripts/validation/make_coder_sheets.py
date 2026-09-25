"""Split the four pair coding sheets into one Excel workbook per coder (B2).

Each coder gets their own copy so the two codings of a record stay independent.
The `openalex_id` column is dropped from the coder copy: an OpenAlex ID leads straight to the
work's topic assignment, which is the thing under test. `record_id` is kept for merging back
against `validation_sample_200_master.json`.

Outputs: data/validation/coder_sheets/coding_sheet_<coder>_records_<start>_<end>.xlsx
"""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
IN_DIR = ROOT / "data" / "validation"
OUT_DIR = IN_DIR / "coder_sheets"
DEADLINE = "Friday 2 October 2026"

PAIRS = [
    ("pair1_bruno_elena", "001_050", ["Bruno Soares", "Elena Gazzea"]),
    ("pair2_stephanie_valentin", "051_100", ["Stephanie Flaman", "Valentin Lucet"]),
    ("pair3_shinichi_losia", "101_150", ["Shinichi Nakagawa", "Losia Lagisz"]),
    ("pair4_marija_eduardo", "151_200", ["Marija Purgar", "Eduardo Santos"]),
]

COLUMNS = [  # (field, header, width)
    ("record_id", "record_id", 9),
    ("year", "year", 7),
    ("venue", "venue", 24),
    ("title", "title", 45),
    ("abstract", "abstract", 80),
    ("authors_and_affiliations", "authors_and_affiliations", 45),
    ("doi", "doi", 20),
    ("is_metaresearch", "Q1 is_metaresearch", 14),
    ("is_canadian", "Q2 is_canadian", 14),
    ("notes", "notes", 35),
]

INSTRUCTIONS = [
    "Canadian Metaresearch Dashboard: validation coding (B2)",
    "",
    "Coder: {coder}",
    "Records: {start}–{end} (50 records, ~1 hour)",
    f"Deadline: {DEADLINE}",
    "",
    "1. Read docs/CODING_GUIDE.md (v1.0) before you start. It is the only rulebook.",
    "2. For every row on the 'Coding' tab, pick YES / NO / UNCLEAR in both Q1 and Q2 (dropdowns).",
    "3. Judge from what is on the row: title, abstract, venue, year, affiliations. You may follow the DOI.",
    "4. Do not look the record up in OpenAlex (or the dashboard): the topic assignment is what we are testing.",
    "5. Do not discuss records with your coding partner until both sheets are submitted.",
    "6. No title and/or no abstract: code what you can; if you cannot decide, UNCLEAR is the right answer.",
    "7. Use 'notes' for anything that made a call hard (one line is enough).",
    "8. Do not add, delete or reorder rows. Return this file as-is (keep the file name).",
]

HEADER_FILL = PatternFill("solid", fgColor="DDE5F0")
ANSWER_FILL = PatternFill("solid", fgColor="FFF6D5")


def slug(name: str) -> str:
    return name.split()[0].lower()


def build(rows: list[dict], coder: str, rng: str) -> Workbook:
    start, end = (int(x) for x in rng.split("_"))
    wb = Workbook()

    info = wb.active
    info.title = "Instructions"
    for i, line in enumerate(INSTRUCTIONS, start=1):
        cell = info.cell(row=i, column=1, value=line.format(coder=coder, start=start, end=end))
        if i == 1:
            cell.font = Font(bold=True, size=14)
    info.column_dimensions["A"].width = 110

    ws = wb.create_sheet("Coding")
    wb.active = 1
    for j, (_, header, width) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=j, value=header)
        c.font = Font(bold=True)
        c.fill = HEADER_FILL
        ws.column_dimensions[c.column_letter].width = width
    for i, r in enumerate(rows, start=2):
        for j, (field, _, _) in enumerate(COLUMNS, start=1):
            value = r.get(field, "")
            if field in ("record_id", "year") and value:
                value = int(value)
            c = ws.cell(row=i, column=j, value=value or None)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if field in ("is_metaresearch", "is_canadian", "notes"):
                c.fill = ANSWER_FILL

    last = len(rows) + 1
    dv = DataValidation(type="list", formula1='"YES,NO,UNCLEAR"', allow_blank=True,
                        showErrorMessage=True, errorTitle="Invalid answer",
                        error="Choose YES, NO or UNCLEAR.")
    ws.add_data_validation(dv)
    dv.add(f"H2:I{last}")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:J{last}"
    return wb


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for pair_id, rng, coders in PAIRS:
        src = IN_DIR / f"coding_sheet_{pair_id}_records_{rng}.csv"
        with open(src, newline="", encoding="utf-8") as f:
            rows = list(csv.DictReader(f))
        assert len(rows) == 50, f"{src.name}: expected 50 records, got {len(rows)}"
        for coder in coders:
            out = OUT_DIR / f"coding_sheet_{slug(coder)}_records_{rng}.xlsx"
            build(rows, coder, rng).save(out)
            print(f"Wrote {coder} ({len(rows)} records) -> {out.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
